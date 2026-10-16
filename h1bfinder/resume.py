from __future__ import annotations
import io, re

ALL_SKILLS = sorted({x.strip().lower() for x in """
python, java, javascript, typescript, c++, c#, go, rust, ruby, php, swift, kotlin, scala, r, matlab, sql, nosql, bash,
react, angular, vue, node.js, nodejs, express, django, flask, fastapi, spring, .net, asp.net, rails, next.js, graphql,
rest, grpc, aws, azure, gcp, kubernetes, docker, terraform, ansible, jenkins, ci/cd, github actions, linux, git,
microservices, kafka, rabbitmq, postgresql, mysql, mongodb, redis, elasticsearch, snowflake, bigquery, redshift,
databricks, spark, hadoop, airflow, dbt, etl, pandas, numpy, scikit-learn, tensorflow, pytorch, keras, nlp, llm, llms,
machine learning, deep learning, computer vision, data science, data engineering, data analysis, data analytics,
statistics, a/b testing, tableau, power bi, looker, excel, mlops, generative ai, prompt engineering, rag,
android, ios, flutter, react native, unity, unreal, embedded, firmware, fpga, verilog, vhdl, rtos,
cybersecurity, penetration testing, siem, iam, devsecops, agile, scrum, jira, product management, roadmap,
salesforce, sap, oracle, servicenow, workday, netsuite, supply chain, logistics, six sigma, lean,
accounting, gaap, financial modeling, valuation, fp&a, audit, risk, compliance, sox, marketing, seo, crm,
cad, solidworks, autocad, ansys, mechanical engineering, electrical engineering,
clinical research, epidemiology, biostatistics, bioinformatics, genomics, ui, ux, figma, design systems
""".replace("\n", " ").split(",") if x.strip()}, key=len, reverse=True)
_RX = {s: re.compile(r"(?<![\w+#.])" + re.escape(s) + r"(?![\w+#]|\.\w)", re.I) for s in ALL_SKILLS}
# 'r', 'go', 'c' style one/two-letter skills are noisy; require case-sensitive standalone tokens
_RX["r"] = re.compile(r"(?<![\w.])R(?![\w.+#])")
_RX["go"] = re.compile(r"(?<![\w.])(Go|Golang)(?![\w.+#])")
_RX["ui"] = re.compile(r"(?<![\w.])UI(?![\w.+#])")
_RX["ux"] = re.compile(r"(?<![\w.])UX(?![\w.+#])")
_RX["ios"] = re.compile(r"(?<![\w.])iOS(?![\w.+#])")


def extract_text(data: bytes, filename: str) -> str:
    fn = filename.lower()
    if fn.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if fn.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


def find_skills(text: str) -> list[str]:
    return sorted(s for s in ALL_SKILLS if _RX[s].search(text or ""))


def years_experience(text: str) -> int:
    ys = [int(m) for m in re.findall(r"(\d{1,2})\+?\s*(?:years|yrs)", text or "", re.I)]
    return max([y for y in ys if y <= 40], default=0)


def build_profile(text: str) -> dict:
    return {"skills": find_skills(text), "years": years_experience(text), "chars": len(text)}
